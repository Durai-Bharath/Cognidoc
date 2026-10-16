"""Offline test helpers: sample documents and a deterministic, prompt-aware MockLLM.

Shared by the test-suite and ``examples/`` so both exercise identical inputs.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pymupdf

from common.llm import MockLLM
from extraction.agents.entity_extraction import fallback_type
from extraction.agents.equation_analysis import text_to_latex

TITLE = "Hydraulic Pump Maintenance Report"
PAGE1 = [
    ("title", TITLE),
    ("h1", "1 Introduction"),
    (
        "p",
        "This report describes the maintenance procedure for the Model X200 hydraulic pump. "
        "The Hydraulic Control Unit (HCU) regulates pressure across the system. Inspection "
        "results are summarised in Table 1. "
        "Section 2 describes the method.",
    ),
    ("h1", "2 Method"),
    (
        "p",
        "Technicians follow ISO 9001 guidelines when servicing the pump. The HCU is "
        "inspected every 500 hours. Pressure is computed with Equation 1.",
    ),
    ("eq", "P = F / A        (1)"),
    ("p", "where P is the pressure, F is the applied force and A is the piston area."),
    ("h2", "2.1 Inspection Results"),
]
TABLE = [
    ["Component", "Interval (h)", "Status"],
    ["Seal", "500", "OK"],
    ["Valve", "1000", "Replaced"],
    ["Filter", "250", "OK"],
]
TABLE_CAPTION = "Table 1: Inspection intervals for pump components."
PAGE2_AFTER_FIGURE = [
    ("cap", "Figure 1: Maintenance process flow for the hydraulic pump."),
    ("h1", "3 Conclusion"),
    (
        "p",
        "The HCU performed within limits during all inspections. Table 1 lists the "
        "inspection intervals and Figure 1 shows the maintenance flow. The method in "
        "Section 2 should be repeated every 500 hours.",
    ),
]
_FONTS = {
    "title": ("hebo", 22),
    "h1": ("hebo", 14),
    "h2": ("hebo", 12),
    "p": ("helv", 10),
    "eq": ("helv", 10),
    "cap": ("helv", 9),
}
MARGIN, WIDTH = 72, 451


def _write_block(page: pymupdf.Page, y: float, kind: str, text: str) -> float:
    font, size = _FONTS[kind]
    lines = max(1, int(len(text) * size * 0.5 / WIDTH) + 1)
    height = lines * size * 1.45 + 4
    rect = pymupdf.Rect(MARGIN, y, MARGIN + WIDTH, y + height)
    page.insert_textbox(rect, text, fontname=font, fontsize=size)
    return y + height + (10 if kind.startswith("h") or kind == "title" else 8)


def _furniture(page: pymupdf.Page, number: int) -> None:
    page.insert_text((MARGIN, 30), "Maintenance Report - Confidential", fontname="helv", fontsize=7)
    page.insert_text((300, 815), str(number), fontname="helv", fontsize=7)


def make_native_pdf(path: str | Path) -> Path:
    """Two-page native PDF with headings, a ruled table, a vector figure, an equation,
    an abbreviation definition and internal citations."""
    path = Path(path)
    doc = pymupdf.open()
    doc.set_metadata({"title": TITLE, "author": "Test Author", "creationDate": "D:20240115120000"})
    page = doc.new_page(width=595, height=842)
    _furniture(page, 1)
    y = 70.0
    for kind, text in PAGE1:
        y = _write_block(page, y, kind, text)
    # ruled table
    col_w, row_h = 120, 20
    for r, row in enumerate(TABLE):
        for c, cell in enumerate(row):
            rect = pymupdf.Rect(
                MARGIN + c * col_w, y + r * row_h, MARGIN + (c + 1) * col_w, y + (r + 1) * row_h
            )
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            page.insert_text(
                (rect.x0 + 4, rect.y1 - 6), cell, fontname="hebo" if r == 0 else "helv", fontsize=9
            )
    y += len(TABLE) * row_h + 8
    _write_block(page, y, "cap", TABLE_CAPTION)

    page2 = doc.new_page(width=595, height=842)
    _furniture(page2, 2)
    # vector flowchart: three boxes joined by arrows
    labels = ["Inspect", "Clean", "Replace"]
    for i, label in enumerate(labels):
        box = pymupdf.Rect(100 + i * 140, 90, 200 + i * 140, 150)
        page2.draw_rect(box, color=(0, 0, 0.6), fill=(0.85, 0.9, 1.0), width=1.2)
        page2.insert_text((box.x0 + 22, box.y0 + 34), label, fontname="helv", fontsize=10)
        if i < len(labels) - 1:
            page2.draw_line((200 + i * 140, 120), (240 + i * 140, 120), width=1.5)
            page2.draw_polyline(
                [(232 + i * 140, 114), (240 + i * 140, 120), (232 + i * 140, 126)], width=1.5
            )
    y = 165.0
    for kind, text in PAGE2_AFTER_FIGURE:
        y = _write_block(page2, y, kind, text)
    doc.save(str(path))
    doc.close()
    return path


def _scan_lines() -> list[tuple[str, float, int, int]]:
    """(text, font scale, thickness, extra gap before) for each rendered scan line."""
    import textwrap

    style = {
        "title": (1.25, 3, 0),
        "h1": (0.95, 2, 26),
        "h2": (0.8, 2, 22),
        "p": (0.62, 1, 12),
        "eq": (0.62, 1, 12),
        "cap": (0.55, 1, 10),
    }
    out: list[tuple[str, float, int, int]] = [("Maintenance Report - Confidential", 0.42, 1, 0)]
    for kind, text in PAGE1:
        scale, thick, gap = style[kind]
        wrapped = textwrap.wrap(text, 78) if kind == "p" else [text]
        for i, line in enumerate(wrapped):
            out.append((line, scale, thick, gap if i == 0 else 0))
    for r, row in enumerate(TABLE):
        out.append(("   ".join(row), 0.62, 2 if r == 0 else 1, 16 if r == 0 else 0))
    out.append((TABLE_CAPTION, 0.55, 1, 10))
    return out


_TABLE_LINES = {"   ".join(row) for row in TABLE}


def make_scanned_image(path: str | Path, *, degrade: bool = True) -> Path:
    """A synthetic scanned page (one OCR line per rendered line). ``degrade`` adds blur,
    noise and washed-out contrast so OCR quality falls below τ_quality."""
    path = Path(path)
    width, height = 1240, 1754  # A4 at 150 dpi
    img = np.full((height, width), 255, dtype=np.uint8)
    y = 90
    for text, scale, thick, gap in _scan_lines():
        y += gap
        line_h = int(32 * scale) + 22
        cells = text.split("   ") if text in _TABLE_LINES else [text]
        for c, cell in enumerate(cells):  # table cells sit at fixed column positions
            cv2.putText(
                img,
                cell,
                (110 + 300 * c, y + int(28 * scale)),
                cv2.FONT_HERSHEY_SIMPLEX,
                scale,
                0,
                thick,
                cv2.LINE_AA,
            )
        y += line_h
    cv2.putText(img, "1", (615, height - 60), cv2.FONT_HERSHEY_SIMPLEX, 0.42, 0, 1, cv2.LINE_AA)
    if degrade:
        f = cv2.GaussianBlur(img.astype(np.float32), (0, 0), 2.2)
        f = f * 0.35 + 110  # washed-out contrast
        rng = np.random.default_rng(7)
        img = np.clip(f + rng.normal(0, 14, f.shape), 0, 255).astype(np.uint8)
    cv2.imwrite(str(path), img)
    return path


def scanned_page_lines() -> list[str]:
    """Text of every rendered scan line, top to bottom, for :class:`MockOCREngine`."""
    return [t for t, *_ in _scan_lines()] + ["1"]


def make_docx(path: str | Path) -> Path:
    import docx

    path = Path(path)
    d = docx.Document()
    d.core_properties.title = "Pump Installation Guide"
    d.core_properties.author = "Docs Team"
    d.add_paragraph("Pump Installation Guide", style="Title")
    d.add_heading("1 Prerequisites", level=1)
    d.add_paragraph(
        "Install the Pump Control Software (PCS) before mounting the unit. "
        "Table 1 lists the required tools."
    )
    t = d.add_table(rows=3, cols=2)
    for r, row in enumerate([["Tool", "Size (mm)"], ["Wrench", "17"], ["Driver", "6"]]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_paragraph("Table 1: Required tools.", style="Caption")
    d.add_heading("2 Installation", level=1)
    d.add_heading("2.1 Mounting", level=2)
    d.add_paragraph("Mount the pump on a level surface and connect it to the PCS.")
    d.save(str(path))
    return path


# ====================================================================== offline LLM


def _section(prompt: str, label: str) -> str:
    m = re.search(rf"{label}:\s*(.*?)(?:\n\n|\n[A-Z][\w ]+:|$)", prompt, re.S)
    return m.group(1).strip() if m else ""


def _summary(prompt: str, _s: str, _i: Sequence[bytes]) -> dict[str, Any]:
    sentences = re.findall(r"^- (.+)$", prompt, re.M)
    summary = " ".join(sentences[:2]) or _section(prompt, "Section title")
    kws = [k.strip() for k in _section(prompt, "Keywords").split(",") if k.strip()]
    return {
        "summary": summary,
        "abstract": sentences[0] if sentences else summary,
        "key_concepts": kws[:4],
        "important_entities": [],
    }


def _table(prompt: str, _s: str, _i: Sequence[bytes]) -> dict[str, Any]:
    caption = _section(prompt, "Caption")
    header = re.search(r"^\| (.+) \|$", prompt, re.M)
    cols = [c.strip() for c in header.group(1).split("|")] if header else []
    return {
        "description": f"{caption} Columns: {', '.join(cols)}.",
        "summary": caption or "Table of values.",
        "keywords": [c.lower() for c in cols][:5],
        "key_attributes": cols[:1],
        "relationships": [],
    }


def _figure(prompt: str, _s: str, _i: Sequence[bytes]) -> dict[str, Any]:
    caption = _section(prompt, "Caption")
    steps = ["Inspect", "Clean", "Replace"]
    return {
        "figure_type": "flowchart",
        "objects": [{"id": s.lower(), "label": s, "kind": "process"} for s in steps],
        "labels": steps,
        "legends": [],
        "annotations": [],
        "connectors": [
            {"source": "inspect", "target": "clean", "label": ""},
            {"source": "clean", "target": "replace", "label": ""},
        ],
        "relationships": ["Inspect precedes Clean", "Clean precedes Replace"],
        "workflow": steps,
        "description": f"{caption} The maintenance flow goes from Inspect to Clean to Replace.",
        "summary": caption,
        "keywords": ["maintenance", "flow"],
    }


def _equation(prompt: str, _s: str, _i: Sequence[bytes]) -> dict[str, Any]:
    eq = _section(prompt, "Equation text")
    context = _section(prompt, "Surrounding text")
    variables = []
    for sym, definition in re.findall(r"\b([A-Z])\b is the ([a-z ]+?)(?:,| and|\.|$)", context):
        if sym in eq:
            variables.append({"symbol": sym, "definition": definition.strip()})
    return {
        "latex": text_to_latex(eq),
        "equation_type": "algebraic",
        "variables": variables,
        "constants": [],
        "operators": ["="],
        "functions": [],
        "constraints": [],
        "description": f"Equation relating {', '.join(v['symbol'] for v in variables)}: {eq}",
        "summary": "Algebraic relation.",
    }


def _entities(prompt: str, _s: str, _i: Sequence[bytes]) -> dict[str, Any]:
    m = re.search(r"NER candidates: (\[.*?\])", prompt)
    names = re.findall(r"'([^']+)'", m.group(1)) if m else []
    return {
        "entities": [
            {
                "name": n,
                "type": fallback_type(n).value,
                "description": f"{n} is mentioned in the section.",
                "attributes": [],
            }
            for n in names[:10]
        ]
    }


def _regenerate(prompt: str, _s: str, _i: Sequence[bytes]) -> dict[str, Any]:
    source = prompt.split("SOURCE:\n", 1)[-1]
    first = re.split(r"(?<=[.!?])\s", source.strip(), maxsplit=1)[0]
    return {"text": first}


def offline_llm(**overrides: Any) -> MockLLM:
    """MockLLM whose answers are derived from the prompt (deterministic, offline)."""
    responses: dict[str, Any] = {
        "summarization": _summary,
        "table_analysis": _table,
        "figure_understanding": _figure,
        "equation_analysis": _equation,
        "entity_extraction": _entities,
        "hallucination_regeneration": _regenerate,
    }
    responses.update(overrides)
    return MockLLM(responses)
